"""Read persisted document content without executing code or Office macros."""
import hashlib
import io
import json
from pathlib import PurePosixPath
import re
import textwrap
import zipfile

from .registry import register_tool
from .documents import validate_spec


MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 200000
MAX_CELLS = 100000
TEXT_FORMATS = {'txt', 'md', 'csv'}


def read_bytes(file):
    with file.file.open('rb') as source:
        data = source.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValueError('Document trop volumineux (10 Mo maximum).')
    return data


def specification_hash(specification):
    return hashlib.sha256(json.dumps(specification, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':')).encode()).hexdigest()


def _cell(value):
    if value is None or value == '':
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    # Dates from imported workbooks are text, never executed or recalculated.
    return value if isinstance(value, (str, int, float, bool)) else str(value)


def _rows(rows):
    normalized = []
    for row in rows:
        values = [_cell(value) for value in row]
        while values and values[-1] is None:
            values.pop()
        normalized.append(values)
    while normalized and not normalized[-1]:
        normalized.pop()
    return normalized


def _pdf_text(text):
    return ' '.join(text.split())


def expected_content(kind, spec):
    """The supported content contract, independent of the file's cached text."""
    validate_document_content(kind, spec)
    title, body = spec.get('title', ''), spec.get('body', '')
    if kind == 'docx':
        tables = []
        for table in spec.get('tables', []):
            rows = table['rows']
            if rows:
                width = max(1, max(map(len, rows)))
                tables.append([[str(value) if value is not None else '' for value in row] +
                               [''] * (width - len(row)) for row in rows])
        return {'blocks': [{'paragraph': value} for value in ([title] if title else []) + body.split('\n')] +
                [{'table': table} for table in tables]}
    if kind == 'xlsx':
        from openpyxl import Workbook
        book = Workbook()
        sheets = []
        for index, table in enumerate(spec.get('tables', []) or [{'rows': [[title], [body]]}]):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = re.sub(r'[\\/*?:\[\]]', '_', table.get('title') or f'Feuille {index + 1}')[:31]
            sheets.append({'title': sheet.title, 'rows': _rows(table['rows'])})
        return {'sheets': sheets}
    if kind == 'pptx':
        return {'slides': [{'title': slide.get('title', ''), 'body': slide.get('body', '')}
                           for slide in spec.get('slides', []) or [{'title': title, 'body': body}]]}
    text = title + '\n\n' + body
    for table in spec.get('tables', []):
        text += '\n\n' + '\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in table['rows'])
    # Match this generator's line wrapping, including long words split across
    # lines, while preserving the distinction between "1 20" and "12 0".
    lines = [line for paragraph in text.split('\n') for line in (textwrap.wrap(paragraph, 90) or [''])]
    return {'text': _pdf_text('\n'.join(lines))}


def validate_document_content(kind, spec):
    validate_spec(kind, spec)
    if kind != 'pptx' and spec.get('slides'):
        raise ValueError('Ce format ne prend pas en charge slides. Place ce contenu dans body ou tables.')
    if kind == 'pptx' and spec.get('tables'):
        raise ValueError('Les tableaux PowerPoint ne sont pas pris en charge. Place le contenu dans slides.')
    if ((kind == 'xlsx' and spec.get('tables')) or (kind == 'pptx' and spec.get('slides'))) and (
        spec.get('title') or spec.get('body')
    ):
        raise ValueError('Place le titre et le texte dans les cellules Excel ou dans les diapositives, pour les conserver.')
    if kind in {'docx', 'pdf'} and any(table.get('title') for table in spec.get('tables', [])):
        raise ValueError('Place les titres de tableaux dans body : ce générateur ne les insère pas automatiquement.')


def inspect_content(kind, data):
    if len(data) > MAX_BYTES:
        raise ValueError('Document trop volumineux (10 Mo maximum).')
    if kind in {'docx', 'xlsx', 'pptx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
                raise ValueError('Document décompressé trop volumineux.')
            if archive.testzip() is not None:
                raise ValueError('Archive du document endommagée.')
    summary = {'format': kind}
    if kind == 'docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        document = Document(io.BytesIO(data))
        blocks, text_parts = [], []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                blocks.append({'paragraph': block.text})
                text_parts.append(block.text)
            else:
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                blocks.append({'table': rows})
                text_parts.append('\n'.join(' | '.join(row) for row in rows))
        content = {'blocks': blocks}
        text = '\n'.join(text_parts)
        summary.update(paragraphs=len(document.paragraphs), tables=len(document.tables))
    elif kind == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        try:
            if len(book.worksheets) > 50:
                raise ValueError('Trop de feuilles (50 maximum).')
            sheets, cells = [], 0
            for sheet in book.worksheets:
                cells += (sheet.max_row or 0) * (sheet.max_column or 0)
                if cells > MAX_CELLS:
                    raise ValueError('Trop de cellules à analyser (100 000 maximum).')
                sheets.append({'title': sheet.title, 'rows': _rows(sheet.iter_rows(values_only=True))})
            content = {'sheets': sheets}
            text = '\n\n'.join(sheet['title'] + '\n' + '\n'.join(
                ' | '.join(str(v) if v is not None else '' for v in row) for row in sheet['rows']) for sheet in sheets)
            summary.update(sheets=[sheet['title'] for sheet in sheets], cells=cells)
        finally:
            book.close()
    elif kind == 'pptx':
        from pptx import Presentation
        deck = Presentation(io.BytesIO(data))
        if len(deck.slides) > 100:
            raise ValueError('Trop de diapositives (100 maximum).')
        slides = []
        for slide in deck.slides:
            title = slide.shapes.title
            body = '\n'.join(shape.text_frame.text for shape in slide.shapes
                             if shape.has_text_frame and (title is None or shape.shape_id != title.shape_id))
            slides.append({'title': title.text if title is not None else '', 'body': body})
        content = {'slides': slides}
        text = '\n\n'.join(slide['title'] + '\n' + slide['body'] for slide in slides)
        summary.update(slides=len(slides))
    elif kind == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 100:
            raise ValueError('PDF chiffré, vide ou de plus de 100 pages.')
        parts, size = [], 0
        for page in reader.pages:
            if not (0 < float(page.mediabox.width) <= 20000 and 0 < float(page.mediabox.height) <= 20000):
                raise ValueError('Dimensions de page invalides.')
            part = page.extract_text() or ''
            size += len(part)
            if size > MAX_TEXT:
                raise ValueError('Texte trop volumineux (200 000 caractères maximum).')
            parts.append(part)
        text = '\n\n'.join(parts)
        content = {'text': _pdf_text(text)}
        summary.update(pages=len(reader.pages))
    elif kind in TEXT_FORMATS:
        text = data.decode('utf-8')
        content = {'text': text}
    else:
        raise ValueError('Formats pris en charge : PDF, DOCX, XLSX, PPTX, TXT, Markdown et CSV.')
    if len(text) > MAX_TEXT or len(json.dumps(content, ensure_ascii=False)) > MAX_TEXT * 2:
        raise ValueError('Contenu trop volumineux à analyser (200 000 caractères maximum).')
    return {'content': content, 'text': text, 'summary': summary,
            'sha256': hashlib.sha256(data).hexdigest()}


def inspect_file(file):
    kind = PurePosixPath(file.original_name.lower()).suffix.lstrip('.')
    return inspect_content(kind, read_bytes(file))


@register_tool(
    name='read_document',
    description=('Lit les octets actuels d’un document PDF, Word, Excel ou PowerPoint, avec tableaux et formules. '
                 'Texte paginé ; utilise next_offset pour lire la suite. Aucun calcul ni contrôle visuel du rendu.'),
    parameters={'type': 'object', 'properties': {
        'file_id': {'type': 'integer'}, 'offset': {'type': 'integer', 'minimum': 0}}, 'required': ['file_id']},
)
def read_document(workspace, file_id, offset=0):
    if not workspace or not isinstance(file_id, int) or isinstance(file_id, bool):
        return {'error': 'Identifiant de fichier invalide.'}
    if not isinstance(offset, int) or isinstance(offset, bool) or offset < 0:
        return {'error': 'Position de lecture invalide.'}
    file = workspace.files.filter(pk=file_id).first()
    if file is None:
        return {'error': 'Document introuvable dans cet espace.'}
    try:
        inspection = inspect_file(file)
        text = inspection['text']
        if not text.strip():
            return {'error': 'Aucun texte exploitable. Pour un document scanné, utilise l’OCR isolé et signale ses limites.'}
        if offset >= len(text):
            return {'error': 'Position située après la fin du document.'}
        end = min(offset + 12000, len(text))
        return {'file_id': file.pk, 'filename': file.original_name, 'content': text[offset:end],
                'offset': offset, 'end': end, 'total_chars': len(text), 'truncated': end < len(text),
                'next_offset': end if end < len(text) else None, 'sha256': inspection['sha256'],
                'summary': inspection['summary'], 'render_verified': False}
    except Exception:
        # Do not expose parser exceptions containing private paths or document data.
        return {'error': 'Document illisible ou dépassant les limites de lecture (10 Mo, 100 pages ou diapositives, 100 000 cellules).'}
