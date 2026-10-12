"""Structured, editable Office/PDF artifacts; no generated code is executed."""
import io
import json
import re
import textwrap
import math
from pathlib import PurePosixPath
from django.core.files.base import ContentFile
from django.db import transaction
from ..models import DocumentArtifact, FileVersion, WorkFile
from ..services import check_quota, extract_text, StorageQuotaError
from .registry import register_tool

MIME = {'docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'pdf':'application/pdf'}


def validate_spec(kind, spec):
    if kind not in MIME or not isinstance(spec,dict) or len(json.dumps(spec,ensure_ascii=False)) > 200000:
        raise ValueError('Format ou contenu invalide (200 000 caractères maximum).')
    if not isinstance(spec.get('title',''),str) or not isinstance(spec.get('body',''),str):
        raise ValueError('Titre et texte invalides.')
    tables = spec.get('tables',[])
    if not isinstance(tables,list) or len(tables)>50:
        raise ValueError('Liste de tableaux invalide (50 maximum).')
    for table in tables:
        if not isinstance(table,dict) or not isinstance(table.get('rows'),list) or len(table['rows']) > 1000:
            raise ValueError('Tableau invalide (1 000 lignes maximum).')
        if not isinstance(table.get('title',''),str):
            raise ValueError('Titre de tableau invalide.')
        for row in table['rows']:
            if not isinstance(row,list) or len(row) > 50:
                raise ValueError('Ligne invalide (50 colonnes maximum).')
            for cell in row:
                if not isinstance(cell,(str,int,float)) and cell is not None:
                    raise ValueError('Cellule invalide.')
                if isinstance(cell,float) and not math.isfinite(cell):
                    raise ValueError('Nombre non fini dans une cellule.')
                if isinstance(cell,str) and cell.startswith('='):
                    allowed = {'SUM','AVERAGE','MIN','MAX','ROUND','IF','COUNT','COUNTA','ABS','SUMIF','COUNTIF'}
                    functions = set(re.findall(r'([A-Za-z][A-Za-z0-9_.]*)\s*\(',cell))
                    if not functions <= allowed or any(v in cell for v in ('[',']','!','|',':/','@',';')):
                        raise ValueError('Formule externe ou non prise en charge.')
    slides = spec.get('slides',[])
    if not isinstance(slides,list) or len(slides) > 100 or any(not isinstance(s,dict) or not isinstance(s.get('title',''),str) or not isinstance(s.get('body',''),str) for s in slides):
        raise ValueError('Diapositives invalides (100 maximum).')


def build_document(kind, spec):
    validate_spec(kind,spec)
    output = io.BytesIO()
    title, body = spec.get('title',''), spec.get('body','')
    if kind == 'docx':
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(11)
        if title:
            doc.add_heading(title,0)
        for paragraph in body.split('\n'):
            doc.add_paragraph(paragraph)
        for item in spec.get('tables',[]):
            rows = item['rows']
            if not rows:
                continue
            table = doc.add_table(rows=0,cols=max(1,max(map(len,rows))))
            table.style = 'Light Shading Accent 1'
            for row in rows:
                for cell,value in zip(table.add_row().cells,row):
                    cell.text = str(value) if value is not None else ''
        doc.save(output)
    elif kind == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.workbook.properties import CalcProperties
        book = Workbook()
        for number,item in enumerate(spec.get('tables',[]) or [{'rows':[[title],[body]]}]):
            sheet = book.active if number == 0 else book.create_sheet()
            sheet.title = re.sub(r'[\\/*?:\[\]]','_',item.get('title') or f'Feuille {number+1}')[:31]
            for row in item['rows']:
                sheet.append(row)
            for cell in sheet[1]:
                cell.font = Font(bold=True,color='FFFFFF')
                cell.fill = PatternFill('solid',fgColor='245B78')
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(60,max(12,max(len(str(c.value or '')) for c in column)+2))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,(int,float)):
                        cell.number_format = '#,##0.00'
            book.calculation = CalcProperties(fullCalcOnLoad=True)
        book.save(output)
    elif kind == 'pptx':
        from pptx import Presentation
        from pptx.util import Pt
        deck = Presentation()
        for item in spec.get('slides',[]) or [{'title':title,'body':body}]:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = item.get('title','')
            slide.placeholders[1].text = item.get('body','')
            for paragraph in slide.placeholders[1].text_frame.paragraphs:
                paragraph.font.size = Pt(20)
        deck.save(output)
    else:
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        text = title + '\n\n' + body
        for table in spec.get('tables',[]):
            text += '\n\n' + '\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in table['rows'])
        try:
            text.encode('cp1252')
        except UnicodeEncodeError as exc:
            raise ValueError('Le PDF prend actuellement en charge les caractères latins. Utilise Word pour ce texte.') from exc
        lines = [line for paragraph in text.split('\n') for line in (textwrap.wrap(paragraph,90) or [''])]
        for offset in range(0,len(lines),48):
            page = writer.add_blank_page(width=595,height=842)
            font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),
                                     NameObject('/BaseFont'):NameObject('/Helvetica'),NameObject('/Encoding'):NameObject('/WinAnsiEncoding')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
            content = bytearray(b'BT /F1 11 Tf 15 TL 45 790 Td\n')
            for line in lines[offset:offset+48]:
                escaped = line.encode('cp1252').replace(b'\\',b'\\\\').replace(b'(',b'\\(').replace(b')',b'\\)')
                content.extend(b'('+escaped+b') Tj T*\n')
            content.extend(b'ET')
            stream = DecodedStreamObject(); stream.set_data(bytes(content))
            page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(output)
    return output.getvalue()


@transaction.atomic
def save_document(workspace, filename, kind, specification):
    return _store_document(workspace,filename,kind,specification,build_document(kind,specification))


@transaction.atomic
def restore_document(file,version):
    with version.binary_file.open('rb') as source:
        data = source.read(10*1024*1024+1)
    return _store_document(file.workspace,file.original_name,file.artifact.kind,version.specification,data)


def _store_document(workspace, filename, kind, specification, data):
    path = PurePosixPath(filename)
    if path.is_absolute() or '..' in path.parts or '\\' in filename or len(filename)>255:
        raise ValueError('Nom de fichier invalide.')
    filename = str(path.with_suffix('.'+kind))
    if len(data) > 10*1024*1024:
        raise ValueError('Document trop volumineux.')
    file = workspace.files.filter(original_name=filename).first()
    check_quota(workspace.user,len(data),new_files=0 if file else 1)
    if file:
        old_name, storage = file.file.name, file.file.storage
        previous = getattr(file,'artifact',None)
        version = FileVersion(file=file,content=file.extracted_text,size_bytes=file.size_bytes,
                              specification=previous.specification if previous else {})
        if file.file:
            with file.file.open('rb') as f:
                version.binary_file.save(path.name,ContentFile(f.read()),save=False)
        version.save()
        if old_name:
            transaction.on_commit(lambda: storage.delete(old_name))
        stale = list(file.versions.order_by('-id').values_list('pk',flat=True)[20:])
        FileVersion.objects.filter(pk__in=stale).delete()
    else:
        file = WorkFile(workspace=workspace,original_name=filename)
    file.file.save(path.name,ContentFile(data),save=False)
    file.size_bytes, file.content_type = len(data), MIME[kind]
    file.extracted_text = extract_text(io.BytesIO(data),MIME[kind],filename)
    file.retrieval_index = {}
    file.save()
    DocumentArtifact.objects.update_or_create(file=file,defaults={'kind':kind,'specification':specification})
    return {'status':'created','file_id':file.pk,'filename':filename,'download_url':f'/bara/{workspace.pk}/files/{file.pk}/download/',
            'preview_url':f'/bara/{workspace.pk}/documents/{file.pk}/','size_bytes':len(data)}


@register_tool(name='create_document',description='Crée ou révise un document Word, Excel, PowerPoint ou PDF avec aperçu et versions. Les formules Excel sont recalculées à l’ouverture.',
    parameters={'type':'object','properties':{'filename':{'type':'string'},'format':{'type':'string','enum':list(MIME)},
        'specification':{'type':'object','properties':{'title':{'type':'string'},'body':{'type':'string'},
            'tables':{'type':'array','items':{'type':'object','properties':{'title':{'type':'string'},'rows':{'type':'array','items':{'type':'array','items':{}}}},'required':['rows']}},
            'slides':{'type':'array','items':{'type':'object','properties':{'title':{'type':'string'},'body':{'type':'string'}}}}}}},
        'required':['filename','format','specification']})
def create_document(workspace, filename, format, specification):
    try:
        return save_document(workspace,filename,format,specification)
    except (ValueError,StorageQuotaError,TypeError,KeyError) as exc:
        return {'error':str(exc)}
