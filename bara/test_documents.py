import io
from django.contrib.auth import get_user_model
from django.test import TestCase
from .models import Workspace
from .tools.documents import save_document, build_document


class DocumentTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('documents-user')
        self.workspace = Workspace.objects.create(user=self.user)

    def test_office_files_reopen_and_formulas_are_preserved(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        word = Document(io.BytesIO(build_document('docx',{'title':'Rapport','body':'Résultat vérifié.'})))
        self.assertIn('Résultat vérifié.',[p.text for p in word.paragraphs])
        ppt = Presentation(io.BytesIO(build_document('pptx',{'slides':[{'title':'Bilan','body':'Deux résultats'}]})))
        self.assertEqual(ppt.slides[0].shapes.title.text,'Bilan')
        book = load_workbook(io.BytesIO(build_document('xlsx',{'tables':[{'rows':[['Article','Montant'],['A',120],['B',80],['Total','=SUM(B2:B3)']]}]})))
        self.assertEqual(book.active['B4'].value,'=SUM(B2:B3)')

    def test_pdf_pagination_and_french_text_extract(self):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(build_document('pdf',{'title':'Étude','body':'Électricité et résultats.\n'*100})))
        self.assertGreater(len(reader.pages),1)
        self.assertIn('Électricité',reader.pages[0].extract_text())

    def test_revision_restore_preview_and_private_access(self):
        result = save_document(self.workspace,'rapport.docx','docx',{'body':'Ancienne version'})
        save_document(self.workspace,'rapport.docx','docx',{'body':'Nouvelle version'})
        file = self.workspace.files.get(pk=result['file_id'])
        self.assertEqual(file.versions.count(),1)
        self.assertTrue(file.versions.first().binary_file)
        self.client.force_login(self.user)
        self.assertContains(self.client.get(result['preview_url']),'Nouvelle version')
        self.client.post(f'/bara/{self.workspace.pk}/files/{file.pk}/history/{file.versions.first().pk}/restore/')
        file.refresh_from_db()
        self.assertIn('Ancienne version',file.extracted_text)
        other = get_user_model().objects.create_user('documents-other')
        self.client.force_login(other)
        self.assertEqual(self.client.get(result['preview_url']).status_code,404)

    def test_external_formulas_and_path_escape_rejected(self):
        with self.assertRaises(ValueError):
            build_document('xlsx',{'tables':[{'rows':[['=WEBSERVICE("https://example.test")']]}]})
        with self.assertRaises(ValueError):
            save_document(self.workspace,'../outside.docx','docx',{'body':'no'})
