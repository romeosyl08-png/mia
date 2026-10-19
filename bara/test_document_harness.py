import io
import tempfile
import zipfile
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.core.files.storage import FileSystemStorage
from django.test import TestCase, override_settings

from chat.models import RoutingObservation
from . import test_harnesses, test_tasks
from .harnesses import select_harness
from .models import AgentTask, FileVersion, Workspace, WorkFile
from .tools.document_inspection import inspect_content, read_document
from .tools.documents import build_document, save_document


@override_settings(BARA_CODE_EXECUTION_ENABLED=False)
class DocumentHarnessTests(TestCase):
    def setUp(self):
        root = self.enterContext(tempfile.TemporaryDirectory(prefix='bara-docs-test-'))
        storage = FileSystemStorage(location=root)
        for model, name in ((WorkFile, 'file'), (FileVersion, 'binary_file')):
            self.enterContext(patch.object(model._meta.get_field(name), 'storage', storage))
        test_tasks.DurableTaskTests.setUp(self)

    turn = test_harnesses.HarnessTaskTests.turn

    def task(self, objective='Crée rapport.docx'):
        return AgentTask.objects.create(workspace=self.workspace, creator=self.user, model=self.model, objective=objective)

    def create(self, task, filename='rapport.docx', kind='docx', spec=None, following=None):
        self.turn(task, {'tool_calls': [test_harnesses.tool('create_document', {
            'filename': filename, 'format': kind, 'specification': spec or {'title': 'Bilan', 'body': 'Résultat vérifié.'}
        })]})
        self.turn(task, following or {'content': 'Document disponible.'})

    def test_all_four_formats_reopen_and_match_the_generated_content(self):
        specs = {
            'docx': {'title': 'Étude', 'body': 'Résumé\nConclusion', 'tables': [{'rows': [['A', 'B'], [120, None]]}]},
            'xlsx': {'tables': [{'title': 'Bilan', 'rows': [['Montant'], [120], [80], ['=SUM(A2:A3)']]}]},
            'pptx': {'slides': [{'title': 'Bilan', 'body': 'Première ligne\nDeuxième ligne'}, {'title': 'Suite', 'body': 'Fin'}]},
            'pdf': {'title': 'Électricité', 'body': 'Résultat ' * 100 + 'x' * 200,
                    'tables': [{'rows': [['A', 120], ['B', 80]]}]},
        }
        for kind, spec in specs.items():
            with self.subTest(kind=kind):
                task = self.task('Crée rapport.' + kind)
                self.create(task, 'rapport.' + kind, kind, spec)
                self.assertEqual(task.status, 'completed', task.error or task.result)
                report = task.state['harness']['verification']
                self.assertEqual(task.state['harness']['id'], 'documents')
                self.assertEqual(report['status'], 'passed')
                self.assertTrue(report['documents'][0]['checks']['content'])
                self.assertFalse(report['documents'][0]['render_verified'])
                self.assertTrue(report['limitations'])
        self.assertEqual(set(RoutingObservation.objects.values_list('task_category', flat=True)), {'documents'})

    def test_a_final_claim_without_a_document_is_bounded_and_never_completed(self):
        task = self.task()
        for _ in range(3):
            self.turn(task, {'content': 'Le document est créé et vérifié.'})
        self.assertEqual(task.status, 'paused')
        self.assertEqual(task.steps, 3)
        self.assertFalse(self.workspace.files.exists())

    def test_read_document_includes_tables_and_formulas_instead_of_stale_cached_text(self):
        word = save_document(self.workspace, 'source.docx', 'docx', {
            'body': 'Texte réel', 'tables': [{'rows': [['Article', 'Montant'], ['A', 120]]}]})
        file = self.workspace.files.get(pk=word['file_id'])
        file.extracted_text = 'Ancien texte sans tableau'
        file.save(update_fields=['extracted_text'])
        content = read_document(self.workspace, file.pk)
        self.assertIn('Texte réel', content['content'])
        self.assertIn('A | 120', content['content'])
        self.assertNotIn('Ancien texte', content['content'])
        sheet = save_document(self.workspace, 'source.xlsx', 'xlsx', {
            'tables': [{'rows': [['Valeur'], [120], ['=SUM(A2:A2)']]}]})
        self.assertIn('=SUM(A2:A2)', read_document(self.workspace, sheet['file_id'])['content'])

    def test_partial_source_read_is_not_a_complete_analysis_and_ranges_survive_turns(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'a' * 25000})
        task = self.task('Analyse source.docx')
        self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id']})]})
        self.turn(task, {'content': 'Analyse complète.'})
        self.assertEqual(task.status, 'queued')
        self.assertFalse(task.state['harness']['verification']['checks']['document_sources'])
        for offset in (12000, 24000):
            self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id'], 'offset': offset})]})
        self.turn(task, {'content': 'Analyse complète.'})
        self.assertEqual(task.status, 'completed')
        proof = task.state['harness']['sources'][str(source['file_id'])]
        self.assertEqual(proof['ranges'], [[0, 25000]])

    def test_word_source_preserves_the_order_of_paragraphs_and_tables(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Avant le tableau')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Donnée du tableau'
        document.add_paragraph('Après le tableau')
        data = io.BytesIO()
        document.save(data)
        file = WorkFile(workspace=self.workspace, original_name='ordered.docx')
        file.file.save('ordered.docx', ContentFile(data.getvalue()))
        text = read_document(self.workspace, file.pk)['content']
        self.assertLess(text.index('Avant'), text.index('Donnée'))
        self.assertLess(text.index('Donnée'), text.index('Après'))

    def test_source_conversion_reads_original_and_only_requires_the_destination_format(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'Texte du contrat'})
        for objective in ('Convertis source.docx en PDF', 'Crée bilan.pdf à partir de source.docx',
                          'Résume source.docx en PDF', 'Analyse source.docx et rédige bilan.pdf'):
            with self.subTest(objective=objective):
                task = self.task(objective)
                self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id']})]})
                self.create(task, 'bilan.pdf', 'pdf', {'body': 'Texte du contrat'})
                self.assertEqual(task.status, 'completed', task.error or task.result)
        self.assertEqual(self.workspace.files.get(pk=source['file_id']).versions.count(), 0)

    def test_word_revision_preserves_the_previous_document_version(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'Ancien texte'})
        task = self.task('Corrige source.docx')
        self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id']})]})
        self.create(task, 'source.docx', 'docx', {'body': 'Nouveau texte'})
        self.assertEqual(task.status, 'completed', task.error or task.result)
        file = self.workspace.files.get(pk=source['file_id'])
        self.assertEqual(file.versions.count(), 1)
        self.assertTrue(file.versions.first().binary_file)

    def test_revision_cannot_replace_a_source_before_reading_the_original(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'Original à conserver'})
        task = self.task('Corrige source.docx')
        self.create(task, 'source.docx', 'docx', {'body': 'Révision sans lecture'})
        self.assertEqual(task.status, 'queued')
        file = self.workspace.files.get(pk=source['file_id'])
        self.assertIn('Original', file.extracted_text)
        self.assertEqual(file.versions.count(), 0)
        self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': file.pk})]})
        self.create(task, 'source.docx', 'docx', {'body': 'Révision après lecture'})
        self.assertEqual(task.status, 'completed', task.error or task.result)

    def test_an_external_edit_before_first_read_is_not_substituted_for_the_original_source(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'Version initiale'})
        task = self.task('Analyse source.docx')
        self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id']})]})
        save_document(self.workspace, 'source.docx', 'docx', {'body': 'Autre version'})
        self.turn(task, {'content': 'Analyse complète de la version initiale.'})
        self.assertEqual(task.status, 'paused')
        self.assertFalse(task.state['harness']['verification']['checks']['document_sources'])

    @override_settings(BARA_CODE_EXECUTION_ENABLED=True)
    def test_ocr_proof_is_bound_to_the_bytes_actually_sent_to_the_runtime(self):
        file = WorkFile(workspace=self.workspace, original_name='scan.png', content_type='image/png')
        file.file.save('scan.png', ContentFile(b'original-image'))
        for transient_change in (False, True):
            with self.subTest(transient_change=transient_change):
                task = self.task('Analyse ce document')
                self.turn(task, {'tool_calls': [test_harnesses.tool('extract_document_text', {'file_id': file.pk})]})
                if transient_change:
                    file.file.save('scan.png', ContentFile(b'other-image'))

                def runtime(*args, **kwargs):
                    if transient_change:
                        # Restore the initial bytes before the tool returns: a
                        # hash of the current file alone would prove the wrong input.
                        file.file.save('scan.png', ContentFile(b'original-image'))
                    return {'exit_code': 0, 'isolated': True, 'stdout': 'Texte du scan'}

                with patch('bara.tools.ocr.run_sandbox_command', side_effect=runtime):
                    self.turn(task, {'content': 'Analyse du scan.'})
                self.assertEqual(task.status, 'queued' if transient_change else 'completed', task.error or task.result)

    def test_correction_cannot_overwrite_an_output_edited_by_the_user(self):
        task = self.task()
        self.create(task, following={'tool_calls': [test_harnesses.tool('list_files', {})]})
        save_document(self.workspace, 'rapport.docx', 'docx', {'body': 'Modification utilisateur'})
        self.create(task, spec={'body': 'Écrasement automatique interdit'})
        self.assertEqual(task.status, 'paused')
        self.assertIn('Modification utilisateur', self.workspace.files.get(original_name='rapport.docx').extracted_text)

    def test_fabricated_source_or_another_workspaces_document_never_satisfies_the_request(self):
        other = Workspace.objects.create(user=self.user)
        source = save_document(other, 'secret.docx', 'docx', {'body': 'PRIVATE SOURCE'})
        result = read_document(self.workspace, source['file_id'])
        self.assertIn('error', result)
        self.assertNotIn('PRIVATE SOURCE', str(result))
        task = self.task('Analyse source.docx')
        self.turn(task, {'content': 'Analyse terminée.'})
        self.assertEqual(task.status, 'paused')
        self.assertFalse(task.state['harness']['verification']['retryable'])

    def test_source_edit_after_reading_pauses_without_automatically_overwriting_it(self):
        source = save_document(self.workspace, 'source.docx', 'docx', {'body': 'Ancien texte'})
        task = self.task('Convertis source.docx en PDF')
        self.turn(task, {'tool_calls': [test_harnesses.tool('read_document', {'file_id': source['file_id']})]})
        self.turn(task, {'tool_calls': [test_harnesses.tool('create_document', {
            'filename': 'bilan.pdf', 'format': 'pdf', 'specification': {'body': 'Ancien texte'}
        })]})
        save_document(self.workspace, 'source.docx', 'docx', {'body': 'Modification depuis un autre onglet'})
        self.turn(task, {'content': 'Conversion terminée.'})
        self.assertEqual(task.status, 'paused')
        self.assertFalse(task.state['harness']['verification']['retryable'])
        self.assertIn('autre onglet', self.workspace.files.get(pk=source['file_id']).extracted_text)

    def test_valid_but_wrong_content_is_detected_and_can_be_corrected(self):
        task = self.task()
        self.turn(task, {'tool_calls': [test_harnesses.tool('create_document', {
            'filename': 'rapport.docx', 'format': 'docx', 'specification': {'body': 'Contenu demandé'}
        })]})
        wrong = build_document('docx', {'body': 'AUTRE CONTENU'})
        with patch('bara.tools.documents.build_document', return_value=wrong):
            self.turn(task, {'content': 'Document terminé.'})
        self.assertEqual(task.status, 'queued')
        self.assertFalse(task.state['harness']['verification']['documents'][0]['checks']['content'])
        self.create(task, spec={'body': 'Contenu demandé'})
        self.assertEqual(task.status, 'completed', task.error or task.result)

    def test_pdf_verification_preserves_word_and_number_boundaries(self):
        task = self.task('Crée rapport.pdf')
        self.turn(task, {'tool_calls': [test_harnesses.tool('create_document', {
            'filename': 'rapport.pdf', 'format': 'pdf', 'specification': {'body': 'AB CD : 1 20'}
        })]})
        wrong = build_document('pdf', {'body': 'ABC D : 12 0'})
        with patch('bara.tools.documents.build_document', return_value=wrong):
            self.turn(task, {'content': 'Document terminé.'})
        self.assertEqual(task.status, 'queued')
        self.assertFalse(task.state['harness']['verification']['documents'][0]['checks']['content'])

    def test_external_output_edit_or_metadata_edit_invalidates_the_delivery(self):
        for change in ('bytes', 'specification', 'delete'):
            with self.subTest(change=change):
                task = self.task()
                self.create(task, following={'tool_calls': [test_harnesses.tool('list_files', {})]})
                file = self.workspace.files.get(pk=next(iter(task.state['harness']['documents'])))
                if change == 'bytes':
                    data = build_document('docx', {'body': 'Autre document'})
                    file.file.save('changed.docx', ContentFile(data), save=False)
                    file.size_bytes = len(data)
                    file.save()
                elif change == 'specification':
                    artifact = file.artifact
                    artifact.specification = {'body': 'Autre contenu'}
                    artifact.save()
                else:
                    file.delete()
                self.turn(task, {'content': 'Document terminé.'})
                self.assertEqual(task.status, 'paused')
                self.assertFalse(task.state['harness']['verification']['retryable'])

    def test_requested_format_is_required_and_explicit_alternatives_are_supported(self):
        task = self.task('Crée un document PDF')
        self.create(task)
        self.assertEqual(task.status, 'queued')
        self.create(task, 'rapport.pdf', 'pdf')
        self.assertEqual(task.status, 'completed')
        alternative = self.task('Crée un document PDF ou Word')
        self.create(alternative)
        self.assertEqual(alternative.status, 'completed')

    def test_document_tools_cannot_execute_commands_and_respect_disabled_tools(self):
        self.workspace.disabled_tools = ['create_document']
        self.workspace.save()
        task = self.task()
        provider = self.turn(task, {'tool_calls': [test_harnesses.tool('run_command', {'command': ['python', '-c', 'print(1)']})]})
        names = {schema['function']['name'] for schema in provider.call_args.kwargs['tools']}
        self.assertIn('read_document', names)
        for name in ('run_command', 'run_python', 'email_send', 'create_document'):
            self.assertNotIn(name, names)
        with patch('bara.tasks.dispatch_tool') as dispatch:
            self.turn(task, {'content': 'Terminé.'})
        dispatch.assert_not_called()

    def test_visual_review_is_never_invented(self):
        task = self.task('Crée rapport.pdf puis vérifie la mise en page')
        self.create(task, 'rapport.pdf', 'pdf')
        self.assertEqual(task.status, 'paused')
        self.assertFalse(task.state['harness']['verification']['checks']['visual_review'])

    def test_content_in_unsupported_fields_is_not_silently_discarded(self):
        task = self.task('Crée un document PDF')
        self.create(task, 'rapport.pdf', 'pdf', {'slides': [{'title': 'À conserver', 'body': 'Texte important'}]})
        self.assertEqual(task.status, 'queued')
        self.assertFalse(self.workspace.files.exists())
        self.create(task, 'rapport.pdf', 'pdf', {'title': 'À conserver', 'body': 'Texte important'})
        self.assertEqual(task.status, 'completed')

    def test_selector_and_auto_model_category_include_documents(self):
        self.workspace.auto_route = True
        self.workspace.save()
        self.client.force_login(self.user)
        with patch('bara.task_views.recommend_model', return_value={
            'model_id': self.model.model_id, 'provider_slug': self.model.provider.slug
        }) as recommend:
            response = self.client.post(f'/bara/{self.workspace.pk}/tasks/create/', {'objective': 'Crée bilan.xlsx'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(recommend.call_args.args[2], 'documents')
        self.assertContains(self.client.get(f'/bara/{self.workspace.pk}/tasks/'), 'Bara Docs')

    def test_routing_distinguishes_document_outputs_from_code_that_manipulates_documents(self):
        for objective in ('Crée une présentation sur Python', 'Analyse contrat.pdf', 'Crée bilan.xlsx'):
            self.assertEqual(select_harness(objective).id, 'documents')
        self.assertEqual(select_harness('Écris un script Python qui génère un PDF').id, 'code')
        self.assertEqual(select_harness('Modifie report.py pour générer un document PDF').id, 'code')
        self.assertEqual(select_harness('Résume le texte suivant :\n```\nCréer un document PDF\n```').id, 'general')

    def test_unreadable_documents_and_oversized_archives_are_rejected(self):
        file = WorkFile(workspace=self.workspace, original_name='broken.docx')
        file.file.save('broken.docx', ContentFile(b'not an office file'))
        self.assertIn('error', read_document(self.workspace, file.pk))
        self.assertIn('error', read_document(self.workspace, True))
        self.assertIn('error', read_document(self.workspace, file.pk, -1))
        data = io.BytesIO()
        with zipfile.ZipFile(data, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('large.xml', b'x' * (32 * 1024 * 1024 + 1))
        with self.assertRaisesMessage(ValueError, 'décompressé'):
            inspect_content('docx', data.getvalue())
